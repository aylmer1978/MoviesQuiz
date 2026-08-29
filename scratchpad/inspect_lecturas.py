from openpyxl import load_workbook

path = r"G:\Mi unidad\ICEC\ICEC 2026\Proyectos\Lecturas2026.xlsx"
wb = load_workbook(path, read_only=True, data_only=True)
print(wb.sheetnames)
ws = wb["PROYECTOS"]
for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True)):
    print(row)
