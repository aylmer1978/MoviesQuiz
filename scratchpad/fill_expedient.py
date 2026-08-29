import re
from openpyxl import load_workbook

base = r"G:\Mi unidad\ICEC\ICEC 2026\Proyectos"
proyectos_path = base + r"\Proyectos_ICEC_2026.xlsx"
lecturas_path = base + r"\Lecturas2026.xlsx"

import unicodedata

def norm(s):
    if s is None:
        return ""
    s = str(s).strip().lower()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^a-z0-9]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s

# load codigo/titulo map
wb1 = load_workbook(proyectos_path, read_only=True, data_only=True)
ws1 = wb1["Proyectos"]
titulo_to_codigo = {}
dupes = set()
for row in ws1.iter_rows(min_row=2, values_only=True):
    codigo, titulo, modalitat = row
    key = norm(titulo)
    if key in titulo_to_codigo and titulo_to_codigo[key] != codigo:
        dupes.add(key)
    titulo_to_codigo[key] = codigo

print("Total titulos cargados:", len(titulo_to_codigo))
if dupes:
    print("ATENCION duplicados de titulo con distinto codigo:", dupes)

# load lecturas, full read-only pass to find header + rows
wb2 = load_workbook(lecturas_path, read_only=True, data_only=True)
ws2 = wb2["PROYECTOS"]

header = None
rows = []
for row in ws2.iter_rows(values_only=True):
    if header is None:
        header = row
        continue
    rows.append(row)

idx_expedient = header.index("EXPEDIENT")
idx_titol = header.index("TÍTOL DEL PROJECTE")
print("Filas de datos:", len(rows))

matched = 0
unmatched = []
for r in rows:
    titulo = r[idx_titol]
    if titulo is None:
        continue
    key = norm(titulo)
    if key in titulo_to_codigo:
        matched += 1
    else:
        unmatched.append(titulo)

print("Matched:", matched)
print("Unmatched count:", len(unmatched))
for u in unmatched:
    print("  NO MATCH:", repr(u))
