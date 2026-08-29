import re
import unicodedata
from openpyxl import load_workbook, Workbook

base = r"G:\Mi unidad\ICEC\ICEC 2026\Proyectos"
proyectos_path = base + r"\Proyectos_ICEC_2026.xlsx"
lecturas_path = base + r"\Lecturas2026.xlsx"
out_path = base + r"\Expedient_Titulo.xlsx"

def norm(s):
    if s is None:
        return ""
    s = str(s).strip().lower()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^a-z0-9]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s

wb1 = load_workbook(proyectos_path, read_only=True, data_only=True)
ws1 = wb1["Proyectos"]
titulo_to_codigo = {}
for row in ws1.iter_rows(min_row=2, values_only=True):
    codigo, titulo, modalitat = row
    titulo_to_codigo[norm(titulo)] = codigo

manual = {
    "desenllaç": "TEC07326000043",
    "un concert a la casa blanca": "TEC073_26_000028",
    "la companyia": "TEC073_26_000097",
}
for k, v in manual.items():
    titulo_to_codigo[norm(k)] = v

wb2 = load_workbook(lecturas_path, read_only=True, data_only=True)
ws2 = wb2["PROYECTOS"]

header = None
rows = []
for row in ws2.iter_rows(values_only=True):
    if header is None:
        header = row
        continue
    if all(v is None for v in row):
        continue
    rows.append(row)

idx_titol = header.index("TÍTOL DEL PROJECTE")

out_rows = []
unmatched = []
for row in rows:
    titulo = row[idx_titol]
    if titulo is None:
        continue
    codigo = titulo_to_codigo.get(norm(titulo))
    if codigo is None:
        unmatched.append(titulo)
    out_rows.append((codigo or "", titulo))

wb_out = Workbook()
ws_out = wb_out.active
ws_out.title = "Expedient"
ws_out.append(["EXPEDIENT", "TÍTOL DEL PROJECTE"])
for codigo, titulo in out_rows:
    ws_out.append([codigo, titulo])

wb_out.save(out_path)
print("Filas escritas:", len(out_rows))
print("Sin match:", unmatched)
print("Guardado en:", out_path)
